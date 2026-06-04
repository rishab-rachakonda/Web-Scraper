from __future__ import annotations

import asyncio
import csv
import io
import json
from datetime import datetime
from pathlib import Path

import aiofiles
import aiosqlite

from core.models import ExportConfig, ScrapedItem, WebhookConfig


class Exporter:
    def __init__(self, config: ExportConfig, job_name: str):
        self._config = config
        self._job_name = job_name
        self._dir = Path(config.output_dir)
        self._dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        self._stem = config.filename or f"{job_name}_{ts}"
        self.written: list[str] = []

    async def export_all(self, items: list[ScrapedItem]) -> list[str]:
        if not items:
            return []
        handlers = {
            "json": self._export_jsonl,        # backward-compat: json => jsonl
            "jsonl": self._export_jsonl,
            "json_pretty": self._export_json_pretty,
            "csv": self._export_csv,
            "xlsx": self._export_xlsx,
            "excel": self._export_xlsx,
            "pdf": self._export_pdf,
            "markdown": self._export_markdown,
            "md": self._export_markdown,
            "html": self._export_html,
            "sqlite": self._export_sqlite,
        }
        tasks = []
        for fmt in dict.fromkeys(self._config.formats):   # dedupe, preserve order
            handler = handlers.get(fmt)
            if handler:
                tasks.append(handler(items))
        paths = await asyncio.gather(*tasks)
        self.written = [str(p) for p in paths if p]
        return self.written

    # ── shared tabular shaping ──────────────────────────────────────────────────

    @staticmethod
    def _stringify(value) -> str:
        if value is None:
            return ""
        if isinstance(value, (dict, list)):
            return json.dumps(value, ensure_ascii=False)
        return str(value)

    def _tabular(self, items: list[ScrapedItem]) -> tuple[list[str], list[dict]]:
        """Data-first columns (then url) with all values stringified — shared by
        the Excel/PDF/Markdown/HTML exporters."""
        data_keys: list[str] = []
        seen: set[str] = set()
        for item in items:
            for k in item.data:
                if k not in seen:
                    seen.add(k)
                    data_keys.append(k)
        fields = data_keys + ["url"]
        rows = []
        for item in items:
            row = {k: self._stringify(item.data.get(k)) for k in data_keys}
            row["url"] = item.url
            rows.append(row)
        return fields, rows

    async def _export_jsonl(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.jsonl"
        async with aiofiles.open(path, "w", encoding="utf-8") as f:
            for item in items:
                row = {
                    "url": item.url,
                    "scraped_at": item.scraped_at.isoformat(),
                    "status_code": item.status_code,
                    "error": item.error,
                    **item.data,
                }
                await f.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
        return path

    async def _export_csv(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.csv"
        meta = {"url", "scraped_at", "status_code", "error"}
        data_keys: list[str] = []
        seen: set[str] = set()
        for item in items:
            for k in item.data:
                if k not in seen:
                    seen.add(k)
                    data_keys.append(k)
        fieldnames = sorted(meta) + data_keys

        buf = io.StringIO()
        writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for item in items:
            row: dict = {
                "url": item.url,
                "scraped_at": item.scraped_at.isoformat(),
                "status_code": item.status_code,
                "error": item.error or "",
                **{
                    k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                    for k, v in item.data.items()
                },
            }
            writer.writerow(row)

        async with aiofiles.open(path, "w", encoding="utf-8", newline="") as f:
            await f.write(buf.getvalue())
        return path

    async def _export_json_pretty(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.json"
        records = [
            {"url": i.url, "scraped_at": i.scraped_at.isoformat(),
             "status_code": i.status_code, "error": i.error, **i.data}
            for i in items
        ]
        async with aiofiles.open(path, "w", encoding="utf-8") as f:
            await f.write(json.dumps(records, ensure_ascii=False, indent=2, default=str))
        return path

    async def _export_xlsx(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.xlsx"
        fields, rows = self._tabular(items)

        def _write():
            from openpyxl import Workbook
            from openpyxl.styles import Font
            wb = Workbook()
            ws = wb.active
            ws.title = self._job_name[:31] or "data"
            ws.append(fields)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in rows:
                ws.append([row[f] for f in fields])
            ws.freeze_panes = "A2"
            for col_idx, f in enumerate(fields, 1):
                width = max(len(f), *(len(str(r[f])) for r in rows[:200])) if rows else len(f)
                ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(width + 2, 60)
            wb.save(path)

        await asyncio.to_thread(_write)
        return path

    async def _export_pdf(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.pdf"
        fields, rows = self._tabular(items)
        max_rows = 1000                      # keep the document sane
        shown, truncated = rows[:max_rows], len(rows) > max_rows

        def l1(v) -> str:
            # Core PDF fonts are Latin-1 only — drop unsupported glyphs.
            return " ".join(str(v).split()).encode("latin-1", "replace").decode("latin-1")

        def _write():
            from fpdf import FPDF
            pdf = FPDF(orientation="L", unit="mm", format="A4")
            pdf.set_auto_page_break(auto=True, margin=12)
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, l1(f"{self._job_name} - {len(rows)} records"), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=7)

            def clip(v: str, n: int = 45) -> str:
                v = l1(v)
                return v[:n] + ("..." if len(v) > n else "")

            with pdf.table(width=pdf.epw, line_height=4.5, text_align="LEFT") as table:
                head = table.row()
                for f in fields:
                    head.cell(clip(f, 25))
                for r in shown:
                    tr = table.row()
                    for f in fields:
                        tr.cell(clip(r[f]))
            if truncated:
                pdf.set_font("Helvetica", "I", 8)
                pdf.cell(0, 6, l1(f"... and {len(rows) - max_rows} more rows (see CSV/JSON for full data)"))
            pdf.output(str(path))

        await asyncio.to_thread(_write)
        return path

    async def _export_markdown(self, items: list[ScrapedItem]) -> Path:
        path = self._dir / f"{self._stem}.md"
        fields, rows = self._tabular(items)

        def esc(v: str) -> str:
            return " ".join(str(v).split()).replace("|", "\\|")

        lines = [f"# {self._job_name} ({len(rows)} records)", ""]
        lines.append("| " + " | ".join(fields) + " |")
        lines.append("| " + " | ".join("---" for _ in fields) + " |")
        for r in rows:
            lines.append("| " + " | ".join(esc(r[f]) for f in fields) + " |")
        async with aiofiles.open(path, "w", encoding="utf-8") as f:
            await f.write("\n".join(lines) + "\n")
        return path

    async def _export_html(self, items: list[ScrapedItem]) -> Path:
        from html import escape
        path = self._dir / f"{self._stem}.html"
        fields, rows = self._tabular(items)
        head = "".join(f"<th>{escape(f)}</th>" for f in fields)
        body = "".join(
            "<tr>" + "".join(f"<td>{escape(str(r[f]))}</td>" for f in fields) + "</tr>"
            for r in rows
        )
        html = (
            "<!doctype html><html><head><meta charset='utf-8'>"
            f"<title>{escape(self._job_name)}</title><style>"
            "body{font-family:system-ui,Arial,sans-serif;margin:24px}"
            "h1{font-size:18px}table{border-collapse:collapse;width:100%}"
            "th,td{border:1px solid #ddd;padding:6px 10px;text-align:left;font-size:14px;vertical-align:top}"
            "th{background:#f4f4f8;position:sticky;top:0}tr:nth-child(even){background:#fafafa}"
            "td{max-width:480px;overflow-wrap:anywhere}</style></head><body>"
            f"<h1>{escape(self._job_name)} <small>({len(rows)} records)</small></h1>"
            f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></body></html>"
        )
        async with aiofiles.open(path, "w", encoding="utf-8") as f:
            await f.write(html)
        return path

    async def _export_sqlite(self, items: list[ScrapedItem]) -> Path:
        db_path = Path(self._config.db_path)
        db_path.parent.mkdir(parents=True, exist_ok=True)

        async with aiosqlite.connect(db_path) as db:
            await db.execute("""
                CREATE TABLE IF NOT EXISTS scraped_items (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    job_name    TEXT,
                    url         TEXT NOT NULL,
                    scraped_at  TEXT,
                    status_code INTEGER,
                    error       TEXT,
                    data        TEXT
                )
            """)
            await db.executemany(
                "INSERT INTO scraped_items "
                "(job_name, url, scraped_at, status_code, error, data) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                [
                    (
                        self._job_name,
                        item.url,
                        item.scraped_at.isoformat(),
                        item.status_code,
                        item.error,
                        json.dumps(item.data, ensure_ascii=False, default=str),
                    )
                    for item in items
                ],
            )
            await db.commit()
        return db_path
